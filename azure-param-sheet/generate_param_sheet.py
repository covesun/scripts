#!/usr/bin/env python3
"""
Azure パラメータシート自動生成スクリプト
使い方: python3 generate_param_sheet.py <azure_params_*.json> [output.xlsx]
"""

import json
import sys
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

# ============================================================
# UI表示名 ↔ APIプロパティ の対応マッピング定義
# ============================================================

VNET_FIELDS = [
    ("名前",                    "name"),
    ("リージョン",              "location"),
    ("アドレス空間",            "addressSpace.addressPrefixes"),
    ("DNSサーバー",             "dhcpOptions.dnsServers"),
    ("DDoS保護標準",            "enableDdosProtection"),
    ("VMの保護",                "enableVmProtection"),
]

SUBNET_FIELDS = [
    ("サブネット名",                        "subnets[*].name"),
    ("アドレス範囲",                        "subnets[*].addressPrefix"),
    ("ネットワークセキュリティグループ",    "subnets[*].networkSecurityGroup.id"),
    ("ルートテーブル",                      "subnets[*].routeTable.id"),
    ("サービスエンドポイント",              "subnets[*].serviceEndpoints[*].service"),
    ("プライベートエンドポイントポリシー",  "subnets[*].privateEndpointNetworkPolicies"),
    ("プライベートリンクポリシー",          "subnets[*].privateLinkServiceNetworkPolicies"),
    ("委任",                                "subnets[*].delegations[*].serviceName"),
]

NSG_RULE_FIELDS = [
    ("ルール名",                "name"),
    ("優先度",                  "priority"),
    ("方向",                    "direction"),
    ("アクセス",                "access"),
    ("プロトコル",              "protocol"),
    ("ソースIPアドレス/範囲",   "sourceAddressPrefix"),
    ("ソースポート範囲",        "sourcePortRange"),
    ("宛先IPアドレス/範囲",     "destinationAddressPrefix"),
    ("宛先ポート範囲",          "destinationPortRange"),
    ("説明",                    "description"),
]

STORAGE_FIELDS = [
    ("名前",                            "name"),
    ("リージョン",                      "location"),
    ("パフォーマンス/SKU",              "sku.name"),
    ("種類",                            "kind"),
    ("アクセス層",                      "accessTier"),
    ("安全な転送が必須",                "enableHttpsTrafficOnly"),
    ("最小TLSバージョン",               "minimumTlsVersion"),
    ("BLOBパブリックアクセス",          "allowBlobPublicAccess"),
    ("共有キーアクセス",                "allowSharedKeyAccess"),
    ("ネットワーク アクセス規則(既定)", "networkRuleSet.defaultAction"),
    ("ネットワーク バイパス",           "networkRuleSet.bypass"),
    ("ファイアウォール IPルール",       "networkRuleSet.ipRules[*].iPAddressOrRange"),
    ("仮想ネットワーク規則",            "networkRuleSet.virtualNetworkRules[*].id"),
    ("大きいファイル共有",              "largeFileSharesState"),
    ("BLOBの論理的な削除",              "properties.deleteRetentionPolicy.enabled"),
    ("論理的な削除の保持日数",          "properties.deleteRetentionPolicy.days"),
    ("階層的名前空間(HNS)",             "isHnsEnabled"),
]

KEYVAULT_FIELDS = [
    ("名前",                            "name"),
    ("リージョン",                      "location"),
    ("価格レベル",                      "properties.sku.name"),
    ("テナントID",                      "properties.tenantId"),
    ("論理的な削除",                    "properties.enableSoftDelete"),
    ("論理的な削除の保持日数",          "properties.softDeleteRetentionInDays"),
    ("消去保護",                        "properties.enablePurgeProtection"),
    ("Azure RBACの承認",                "properties.enableRbacAuthorization"),
    ("パブリックネットワークアクセス",  "properties.publicNetworkAccess"),
    ("ネットワーク アクセス規則(既定)", "properties.networkAcls.defaultAction"),
    ("ネットワーク バイパス",           "properties.networkAcls.bypass"),
    ("ファイアウォール IPルール",       "properties.networkAcls.ipRules[*].value"),
    ("仮想ネットワーク規則",            "properties.networkAcls.virtualNetworkRules[*].id"),
]

CAE_FIELDS = [
    ("名前",                            "name"),
    ("リージョン",                      "location"),
    ("インフラサブネット",              "properties.vnetConfiguration.infrastructureSubnetId"),
    ("内部のみ(VNet統合)",              "properties.vnetConfiguration.internal"),
    ("DockerブリッジCIDR",              "properties.vnetConfiguration.dockerBridgeCidr"),
    ("プラットフォーム予約CIDR",        "properties.vnetConfiguration.platformReservedCidr"),
    ("プラットフォーム予約DNS IP",      "properties.vnetConfiguration.platformReservedDnsIP"),
    ("Log AnalyticsワークスペースID",   "properties.appLogsConfiguration.logAnalyticsConfiguration.customerId"),
    ("ゾーン冗長",                      "properties.zoneRedundant"),
]

CA_FIELDS = [
    ("名前",                        "name"),
    ("リージョン",                  "location"),
    ("環境",                        "properties.environmentId"),
    ("イングレス(外部公開)",        "properties.configuration.ingress.external"),
    ("ターゲットポート",            "properties.configuration.ingress.targetPort"),
    ("トランスポート",              "properties.configuration.ingress.transport"),
    ("最小レプリカ数",              "properties.template.scale.minReplicas"),
    ("最大レプリカ数",              "properties.template.scale.maxReplicas"),
    ("コンテナイメージ",            "properties.template.containers[*].image"),
    ("CPU",                         "properties.template.containers[*].resources.cpu"),
    ("メモリ",                      "properties.template.containers[*].resources.memory"),
    ("マネージドID",                "identity.type"),
    ("コンテナレジストリ",          "properties.configuration.registries[*].server"),
]

AGW_TOP_FIELDS = [
    ("名前",                          "name"),
    ("リージョン",                    "location"),
    ("SKU名",                         "sku.name"),
    ("SKUティア",                     "sku.tier"),
    ("容量",                          "sku.capacity"),
    ("最小容量(自動スケール)",        "autoscaleConfiguration.minCapacity"),
    ("最大容量(自動スケール)",        "autoscaleConfiguration.maxCapacity"),
    ("WAF有効",                       "webApplicationFirewallConfiguration.enabled"),
    ("WAFモード",                     "webApplicationFirewallConfiguration.firewallMode"),
    ("WAFルールセットタイプ",         "webApplicationFirewallConfiguration.ruleSetType"),
    ("WAFルールセットバージョン",     "webApplicationFirewallConfiguration.ruleSetVersion"),
    ("フロントエンドポート",          "frontendPorts[*].port"),
]

AGW_POOL_FIELDS = [
    ("プール名",                      "name"),
    ("IPアドレス",                    "backendAddresses[*].ipAddress"),
    ("FQDN",                          "backendAddresses[*].fqdn"),
]

AGW_HTTP_SETTINGS_FIELDS = [
    ("設定名",                        "name"),
    ("ポート",                        "port"),
    ("プロトコル",                    "protocol"),
    ("クッキーアフィニティ",          "cookieBasedAffinity"),
    ("タイムアウト(秒)",              "requestTimeout"),
    ("ホスト名(バックエンドから)",    "pickHostNameFromBackendAddress"),
    ("カスタムプローブ",              "probe.id"),
]

AGW_LISTENER_FIELDS = [
    ("リスナー名",                    "name"),
    ("プロトコル",                    "protocol"),
    ("フロントエンドポート",          "frontendPort.id"),
    ("ホスト名",                      "hostName"),
    ("SSL証明書",                     "sslCertificate.id"),
]

AGW_RULE_FIELDS = [
    ("ルール名",                      "name"),
    ("優先度",                        "priority"),
    ("ルールタイプ",                  "ruleType"),
    ("リスナー",                      "httpListener.id"),
    ("バックエンドプール",            "backendAddressPool.id"),
    ("バックエンドHTTP設定",          "backendHttpSettings.id"),
    ("URLパスマップ",                 "urlPathMap.id"),
    ("リダイレクト設定",              "redirectConfiguration.id"),
]

FIREWALL_FIELDS = [
    ("名前",                          "name"),
    ("リージョン",                    "location"),
    ("SKU名",                         "sku.name"),
    ("SKUティア",                     "sku.tier"),
    ("ファイアウォールポリシー",      "firewallPolicy.id"),
    ("脅威インテリジェンスモード",    "threatIntelMode"),
    ("IPアドレス設定名",              "ipConfigurations[*].name"),
    ("サブネット",                    "ipConfigurations[*].subnet.id"),
    ("パブリックIP",                  "ipConfigurations[*].publicIPAddress.id"),
    ("プライベートIPアドレス",        "ipConfigurations[*].privateIPAddress"),
    ("管理IP設定名",                  "managementIpConfiguration.name"),
    ("Availability Zones",            "zones"),
]

# ============================================================
# スタイル定義
# ============================================================

COLOR_HEADER_SECTION = "1F4E79"   # 濃い青（シートタイトル）
COLOR_HEADER_COL     = "2E75B6"   # 青（列ヘッダー）
COLOR_SUBHEADER      = "D6E4F0"   # 薄い青（サブヘッダー行）
COLOR_ODD_ROW        = "FFFFFF"
COLOR_EVEN_ROW       = "F2F7FB"
COLOR_RESOURCE_ROW   = "E2EFDA"   # 薄い緑（リソース区切り行）

def cell_font(bold=False, color="000000", size=10):
    return Font(name="Arial", bold=bold, color=color, size=size)

def cell_fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def thin_border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)

def set_header(ws, row, col, value, bg=COLOR_HEADER_COL, fg="FFFFFF", bold=True, size=10):
    c = ws.cell(row=row, column=col, value=value)
    c.font = cell_font(bold=bold, color=fg, size=size)
    c.fill = cell_fill(bg)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = thin_border()
    return c

def set_data(ws, row, col, value, bg=COLOR_ODD_ROW):
    c = ws.cell(row=row, column=col, value=value)
    c.font = cell_font()
    c.fill = cell_fill(bg)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    c.border = thin_border()
    return c

def resolve(obj, path):
    """ネストされたパスを解決してわかりやすい文字列を返す"""
    if obj is None:
        return ""
    
    # 配列アクセスを含むパス
    if "[*]" in path:
        parts = path.split("[*].")
        base_path = parts[0]
        rest = parts[1] if len(parts) > 1 else None
        
        val = resolve(obj, base_path)
        if isinstance(val, list):
            if rest:
                items = [resolve(i, rest) for i in val if i]
                items = [str(i) for i in items if i != ""]
            else:
                items = [str(i) for i in val if i is not None]
            return ", ".join(items) if items else ""
        return ""
    
    # 通常のドット区切りパス
    keys = path.split(".")
    current = obj
    for k in keys:
        if current is None:
            return ""
        if isinstance(current, dict):
            current = current.get(k)
        elif isinstance(current, list):
            current = [item.get(k) for item in current if isinstance(item, dict)]
        else:
            return ""
    
    if isinstance(current, list):
        items = [str(i) for i in current if i is not None]
        return ", ".join(items) if items else ""
    if isinstance(current, bool):
        return "有効" if current else "無効"
    if current is None:
        return ""
    return str(current)

# ============================================================
# シート生成関数
# ============================================================

def write_section_title(ws, row, title, ncols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=title)
    c.font = cell_font(bold=True, color="FFFFFF", size=12)
    c.fill = cell_fill(COLOR_HEADER_SECTION)
    c.alignment = Alignment(horizontal="left", vertical="center")
    c.border = thin_border()
    ws.row_dimensions[row].height = 22

def create_vnet_sheet(wb, vnets):
    ws = wb.create_sheet("VNet・Subnet")
    ws.freeze_panes = "A3"

    # ---- VNet セクション ----
    write_section_title(ws, 1, "■ Virtual Network", 10)
    headers = ["#", "UI項目名", "設定値", "APIプロパティパス", "備考"]
    col_widths = [5, 30, 40, 45, 25]
    for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
        set_header(ws, 2, ci, h)
        ws.column_dimensions[get_column_letter(ci)].width = w

    row = 3
    for vnet in vnets:
        # リソース名行
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        c = ws.cell(row=row, column=1, value=f"VNet: {vnet.get('name', '')}")
        c.font = cell_font(bold=True, size=10)
        c.fill = cell_fill(COLOR_RESOURCE_ROW)
        c.border = thin_border()
        ws.row_dimensions[row].height = 18
        row += 1

        for idx, (ui_name, api_path) in enumerate(VNET_FIELDS):
            bg = COLOR_ODD_ROW if idx % 2 == 0 else COLOR_EVEN_ROW
            set_data(ws, row, 1, idx + 1, bg)
            set_data(ws, row, 2, ui_name, bg)
            set_data(ws, row, 3, resolve(vnet, api_path), bg)
            set_data(ws, row, 4, api_path, bg)
            ws.cell(row=row, column=4).font = Font(name="Courier New", size=9, color="555555")
            set_data(ws, row, 5, "", bg)
            row += 1

        row += 1  # 空行

    # ---- Subnet セクション ----
    write_section_title(ws, row, "■ Subnet", 10)
    row += 1
    headers_sub = ["#", "VNet名", "UI項目名", "設定値", "APIプロパティパス", "備考"]
    col_widths_sub = [5, 25, 35, 40, 45, 25]
    for ci, (h, w) in enumerate(zip(headers_sub, col_widths_sub), 1):
        set_header(ws, row, ci, h)
        ws.column_dimensions[get_column_letter(ci)].width = max(
            ws.column_dimensions[get_column_letter(ci)].width, w)
    row += 1

    for vnet in vnets:
        subnets = vnet.get("subnets", [])
        for subnet in subnets:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            c = ws.cell(row=row, column=1,
                        value=f"Subnet: {vnet.get('name','')} / {subnet.get('name','')}")
            c.font = cell_font(bold=True, size=10)
            c.fill = cell_fill(COLOR_RESOURCE_ROW)
            c.border = thin_border()
            ws.row_dimensions[row].height = 18
            row += 1

            subnet_ui = [
                ("サブネット名",                       "name"),
                ("アドレス範囲",                       "addressPrefix"),
                ("ネットワークセキュリティグループ",   "networkSecurityGroup.id"),
                ("ルートテーブル",                     "routeTable.id"),
                ("サービスエンドポイント",             "serviceEndpoints[*].service"),
                ("プライベートエンドポイントポリシー", "privateEndpointNetworkPolicies"),
                ("プライベートリンクポリシー",         "privateLinkServiceNetworkPolicies"),
                ("委任",                               "delegations[*].serviceName"),
            ]
            for idx, (ui_name, api_path) in enumerate(subnet_ui):
                bg = COLOR_ODD_ROW if idx % 2 == 0 else COLOR_EVEN_ROW
                set_data(ws, row, 1, idx + 1, bg)
                set_data(ws, row, 2, vnet.get("name", ""), bg)
                set_data(ws, row, 3, ui_name, bg)
                set_data(ws, row, 4, resolve(subnet, api_path), bg)
                set_data(ws, row, 5, api_path, bg)
                ws.cell(row=row, column=5).font = Font(name="Courier New", size=9, color="555555")
                set_data(ws, row, 6, "", bg)
                row += 1
            row += 1

    ws.row_dimensions[1].height = 22
    return ws


def create_nsg_sheet(wb, nsgs):
    ws = wb.create_sheet("NSG")
    ws.freeze_panes = "A3"
    write_section_title(ws, 1, "■ Network Security Group", 8)

    headers = ["#", "NSG名", "ルール名", "優先度", "方向", "アクセス",
               "プロトコル", "ソースIP", "ソースPort", "宛先IP", "宛先Port", "説明", "備考"]
    col_widths = [5, 25, 25, 10, 10, 10, 12, 25, 15, 25, 15, 30, 25]
    for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
        set_header(ws, 2, ci, h)
        ws.column_dimensions[get_column_letter(ci)].width = w

    row = 3
    for nsg in nsgs:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        c = ws.cell(row=row, column=1, value=f"NSG: {nsg.get('name', '')}")
        c.font = cell_font(bold=True, size=10)
        c.fill = cell_fill(COLOR_RESOURCE_ROW)
        c.border = thin_border()
        ws.row_dimensions[row].height = 18
        row += 1

        rules = nsg.get("securityRules", [])
        for idx, rule in enumerate(rules):
            bg = COLOR_ODD_ROW if idx % 2 == 0 else COLOR_EVEN_ROW
            vals = [
                idx + 1,
                nsg.get("name", ""),
                rule.get("name", ""),
                rule.get("priority", ""),
                rule.get("direction", ""),
                rule.get("access", ""),
                rule.get("protocol", ""),
                rule.get("sourceAddressPrefix", ""),
                rule.get("sourcePortRange", ""),
                rule.get("destinationAddressPrefix", ""),
                rule.get("destinationPortRange", ""),
                rule.get("description", ""),
                "",
            ]
            for ci, v in enumerate(vals, 1):
                set_data(ws, row, ci, v, bg)
            row += 1
        row += 1

    return ws


def create_generic_sheet(wb, sheet_name, title, resources, fields):
    ws = wb.create_sheet(sheet_name)
    ws.freeze_panes = "A3"
    write_section_title(ws, 1, title, 5)

    headers = ["#", "UI項目名", "設定値", "APIプロパティパス", "備考"]
    col_widths = [5, 35, 45, 50, 25]
    for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
        set_header(ws, 2, ci, h)
        ws.column_dimensions[get_column_letter(ci)].width = w

    row = 3
    for resource in resources:
        name = resource.get("name", resource.get("id", "unknown"))
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        c = ws.cell(row=row, column=1, value=f"リソース: {name}")
        c.font = cell_font(bold=True, size=10)
        c.fill = cell_fill(COLOR_RESOURCE_ROW)
        c.border = thin_border()
        ws.row_dimensions[row].height = 18
        row += 1

        for idx, (ui_name, api_path) in enumerate(fields):
            bg = COLOR_ODD_ROW if idx % 2 == 0 else COLOR_EVEN_ROW
            set_data(ws, row, 1, idx + 1, bg)
            set_data(ws, row, 2, ui_name, bg)
            set_data(ws, row, 3, resolve(resource, api_path), bg)
            set_data(ws, row, 4, api_path, bg)
            ws.cell(row=row, column=4).font = Font(name="Courier New", size=9, color="555555")
            set_data(ws, row, 5, "", bg)
            row += 1
        row += 1

    return ws


def create_agw_sheet(wb, agws):
    """Application Gateway専用シート（サブセクション構成）"""
    ws = wb.create_sheet("AppGateway")
    ws.freeze_panes = "A3"
    headers = ["#", "UI項目名", "設定値", "APIプロパティパス", "備考"]
    col_widths = [5, 35, 45, 50, 25]

    def _section(title, row):
        write_section_title(ws, row, title, 5)
        row += 1
        for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
            set_header(ws, row, ci, h)
            ws.column_dimensions[get_column_letter(ci)].width = max(
                ws.column_dimensions[get_column_letter(ci)].width, w)
        return row + 1

    def _resource_header(label, row):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        c = ws.cell(row=row, column=1, value=label)
        c.font = cell_font(bold=True, size=10)
        c.fill = cell_fill(COLOR_RESOURCE_ROW)
        c.border = thin_border()
        ws.row_dimensions[row].height = 18
        return row + 1

    def _kvp_rows(obj, fields, row):
        for idx, (ui_name, api_path) in enumerate(fields):
            bg = COLOR_ODD_ROW if idx % 2 == 0 else COLOR_EVEN_ROW
            set_data(ws, row, 1, idx + 1, bg)
            set_data(ws, row, 2, ui_name, bg)
            set_data(ws, row, 3, resolve(obj, api_path), bg)
            set_data(ws, row, 4, api_path, bg)
            ws.cell(row=row, column=4).font = Font(name="Courier New", size=9, color="555555")
            set_data(ws, row, 5, "", bg)
            row += 1
        return row + 1  # 空行

    # ---- AGW本体 ----
    row = _section("■ Application Gateway", 1)
    for agw in agws:
        row = _resource_header(f"AGW: {agw.get('name', '')}", row)
        row = _kvp_rows(agw, AGW_TOP_FIELDS, row)

    # ---- バックエンドプール ----
    row = _section("■ バックエンドプール", row)
    for agw in agws:
        agw_name = agw.get("name", "")
        for pool in agw.get("backendAddressPools", []):
            row = _resource_header(f"バックエンドプール: {agw_name} / {pool.get('name','')}", row)
            row = _kvp_rows(pool, AGW_POOL_FIELDS, row)

    # ---- バックエンドHTTP設定 ----
    row = _section("■ バックエンドHTTP設定", row)
    for agw in agws:
        agw_name = agw.get("name", "")
        for setting in agw.get("backendHttpSettingsCollection", []):
            obj = setting.get("properties", setting)
            obj.setdefault("name", setting.get("name", ""))
            row = _resource_header(f"HTTP設定: {agw_name} / {setting.get('name','')}", row)
            row = _kvp_rows(obj, AGW_HTTP_SETTINGS_FIELDS, row)

    # ---- HTTPリスナー ----
    row = _section("■ HTTPリスナー", row)
    for agw in agws:
        agw_name = agw.get("name", "")
        for listener in agw.get("httpListeners", []):
            obj = listener.get("properties", listener)
            obj.setdefault("name", listener.get("name", ""))
            row = _resource_header(f"リスナー: {agw_name} / {listener.get('name','')}", row)
            row = _kvp_rows(obj, AGW_LISTENER_FIELDS, row)

    # ---- ルーティングルール ----
    row = _section("■ ルーティングルール", row)
    for agw in agws:
        agw_name = agw.get("name", "")
        for rule in agw.get("requestRoutingRules", []):
            obj = rule.get("properties", rule)
            obj.setdefault("name", rule.get("name", ""))
            row = _resource_header(f"ルーティングルール: {agw_name} / {rule.get('name','')}", row)
            row = _kvp_rows(obj, AGW_RULE_FIELDS, row)

    return ws


def create_summary_sheet(wb, data):
    ws = wb.create_sheet("サマリー", 0)
    write_section_title(ws, 1, "Azure パラメータシート", 4)

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 30

    set_header(ws, 2, 1, "リソース種別")
    set_header(ws, 2, 2, "件数")
    set_header(ws, 2, 3, "シート名")
    set_header(ws, 2, 4, "備考")

    summary = [
        ("VNet / Subnet",           len(data.get("vnets", [])),                "VNet・Subnet"),
        ("NSG",                     len(data.get("nsgs", [])),                 "NSG"),
        ("Storage Account",         len(data.get("storageAccounts", [])),      "StorageAccount"),
        ("Key Vault",               len(data.get("keyVaults", [])),            "KeyVault"),
        ("Container Apps Env",      len(data.get("containerAppEnvs", [])),     "ContainerAppsEnv"),
        ("Container Apps",          len(data.get("containerApps", [])),        "ContainerApps"),
        ("Application Gateway",     len(data.get("applicationGateways", [])),  "AppGateway"),
        ("Azure Firewall",          len(data.get("firewalls", [])),            "Firewall"),
    ]
    for idx, (rtype, cnt, sname) in enumerate(summary, 3):
        bg = COLOR_ODD_ROW if idx % 2 == 0 else COLOR_EVEN_ROW
        set_data(ws, idx, 1, rtype, bg)
        set_data(ws, idx, 2, cnt, bg)
        set_data(ws, idx, 3, sname, bg)
        set_data(ws, idx, 4, "", bg)

    ws.cell(row=12, column=1, value=f"対象RG: {data.get('resourceGroup', '')}")
    ws.cell(row=12, column=1).font = cell_font(bold=True)

    return ws


# ============================================================
# メイン
# ============================================================

def main():
    if len(sys.argv) < 2:
        print("Usage: python3 generate_param_sheet.py <azure_params.json> [output.xlsx]")
        sys.exit(1)

    input_file = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) > 2 else \
        input_file.replace(".json", "_paramsheet.xlsx")

    with open(input_file, encoding="utf-8") as f:
        data = json.load(f)

    wb = Workbook()
    # デフォルトシート削除
    wb.remove(wb.active)

    create_summary_sheet(wb, data)
    create_vnet_sheet(wb, data.get("vnets", []))
    create_nsg_sheet(wb, data.get("nsgs", []))
    create_generic_sheet(wb, "StorageAccount", "■ Storage Account",
                         data.get("storageAccounts", []), STORAGE_FIELDS)
    create_generic_sheet(wb, "KeyVault", "■ Key Vault",
                         data.get("keyVaults", []), KEYVAULT_FIELDS)
    create_generic_sheet(wb, "ContainerAppsEnv", "■ Container Apps Environment",
                         data.get("containerAppEnvs", []), CAE_FIELDS)
    create_generic_sheet(wb, "ContainerApps", "■ Container Apps",
                         data.get("containerApps", []), CA_FIELDS)
    create_agw_sheet(wb, data.get("applicationGateways", []))
    create_generic_sheet(wb, "Firewall", "■ Azure Firewall",
                         data.get("firewalls", []), FIREWALL_FIELDS)

    wb.save(output_file)
    print(f"パラメータシート生成完了: {output_file}")


if __name__ == "__main__":
    main()
