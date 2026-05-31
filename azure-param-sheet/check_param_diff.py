#!/usr/bin/env python3
"""
Azure パラメータ突合チェックスクリプト
使い方: python3 check_param_diff.py <design.xlsx> <azure_params.json> [output.xlsx]

出力列構成: # / UI項目名 / 設計値 / 実環境値 / 判定 / APIプロパティパス / 備考
差分セル: 赤ハイライト
"""

import json
import sys
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ============================================================
# フィールド定義（generate_param_sheet.py と同一）
# ============================================================

VNET_FIELDS = [
    ("名前",         "name"),
    ("リージョン",   "location"),
    ("アドレス空間", "addressSpace.addressPrefixes"),
    ("DNSサーバー",  "dhcpOptions.dnsServers"),
    ("DDoS保護標準", "enableDdosProtection"),
    ("VMの保護",     "enableVmProtection"),
]

SUBNET_FIELDS = [
    ("サブネット名",                       "name"),
    ("アドレス範囲",                       "addressPrefix"),
    ("ネットワークセキュリティグループ",   "networkSecurityGroup.id"),
    ("ルートテーブル",                     "routeTable.id"),
    ("サービスエンドポイント",             "serviceEndpoints[*].service"),
    ("プライベートエンドポイントポリシー", "privateEndpointNetworkPolicies"),
    ("プライベートリンクポリシー",         "privateLinkServiceNetworkPolicies"),
    ("委任",                               "delegations[*].serviceName"),
]

NSG_RULE_COLS = [
    "name", "priority", "direction", "access", "protocol",
    "sourceAddressPrefix", "sourcePortRange",
    "destinationAddressPrefix", "destinationPortRange", "description",
]

STORAGE_FIELDS = [
    ("名前",                            "name"),
    ("リージョン",                      "location"),
    ("パフォーマンス/SKU",              "sku.name"),
    ("種類",                            "kind"),
    ("アクセス層",                      "accessTier"),
    ("安全な転送が必須",                "enableHttpsTrafficOnly"),
    ("最小TLSバージョン",               "minimumTlsVersion"),
    ("BLOBパブリックアクセス",          "allowBlobPublicAccess"),
    ("共有キーアクセス",                "allowSharedKeyAccess"),
    ("ネットワーク アクセス規則(既定)", "networkRuleSet.defaultAction"),
    ("ネットワーク バイパス",           "networkRuleSet.bypass"),
    ("ファイアウォール IPルール",       "networkRuleSet.ipRules[*].iPAddressOrRange"),
    ("仮想ネットワーク規則",            "networkRuleSet.virtualNetworkRules[*].id"),
    ("大きいファイル共有",              "largeFileSharesState"),
    ("BLOBの論理的な削除",              "properties.deleteRetentionPolicy.enabled"),
    ("論理的な削除の保持日数",          "properties.deleteRetentionPolicy.days"),
    ("階層的名前空間(HNS)",             "isHnsEnabled"),
]

KEYVAULT_FIELDS = [
    ("名前",                            "name"),
    ("リージョン",                      "location"),
    ("価格レベル",                      "properties.sku.name"),
    ("テナントID",                      "properties.tenantId"),
    ("論理的な削除",                    "properties.enableSoftDelete"),
    ("論理的な削除の保持日数",          "properties.softDeleteRetentionInDays"),
    ("消去保護",                        "properties.enablePurgeProtection"),
    ("Azure RBACの承認",                "properties.enableRbacAuthorization"),
    ("パブリックネットワークアクセス",  "properties.publicNetworkAccess"),
    ("ネットワーク アクセス規則(既定)", "properties.networkAcls.defaultAction"),
    ("ネットワーク バイパス",           "properties.networkAcls.bypass"),
    ("ファイアウォール IPルール",       "properties.networkAcls.ipRules[*].value"),
    ("仮想ネットワーク規則",            "properties.networkAcls.virtualNetworkRules[*].id"),
]

CAE_FIELDS = [
    ("名前",                          "name"),
    ("リージョン",                    "location"),
    ("インフラサブネット",            "properties.vnetConfiguration.infrastructureSubnetId"),
    ("内部のみ(VNet統合)",            "properties.vnetConfiguration.internal"),
    ("DockerブリッジCIDR",            "properties.vnetConfiguration.dockerBridgeCidr"),
    ("プラットフォーム予約CIDR",      "properties.vnetConfiguration.platformReservedCidr"),
    ("プラットフォーム予約DNS IP",    "properties.vnetConfiguration.platformReservedDnsIP"),
    ("Log AnalyticsワークスペースID", "properties.appLogsConfiguration.logAnalyticsConfiguration.customerId"),
    ("ゾーン冗長",                    "properties.zoneRedundant"),
]

CA_FIELDS = [
    ("名前",               "name"),
    ("リージョン",         "location"),
    ("環境",               "properties.environmentId"),
    ("イングレス(外部公開)", "properties.configuration.ingress.external"),
    ("ターゲットポート",   "properties.configuration.ingress.targetPort"),
    ("トランスポート",     "properties.configuration.ingress.transport"),
    ("最小レプリカ数",     "properties.template.scale.minReplicas"),
    ("最大レプリカ数",     "properties.template.scale.maxReplicas"),
    ("コンテナイメージ",   "properties.template.containers[*].image"),
    ("CPU",                "properties.template.containers[*].resources.cpu"),
    ("メモリ",             "properties.template.containers[*].resources.memory"),
    ("マネージドID",       "identity.type"),
    ("コンテナレジストリ", "properties.configuration.registries[*].server"),
]

AGW_TOP_FIELDS = [
    ("名前",                          "name"),
    ("リージョン",                    "location"),
    ("SKU名",                         "sku.name"),
    ("SKUティア",                     "sku.tier"),
    ("容量",                          "sku.capacity"),
    ("最小容量(自動スケール)",        "autoscaleConfiguration.minCapacity"),
    ("最大容量(自動スケール)",        "autoscaleConfiguration.maxCapacity"),
    ("WAF有効",                       "webApplicationFirewallConfiguration.enabled"),
    ("WAFモード",                     "webApplicationFirewallConfiguration.firewallMode"),
    ("WAFルールセットタイプ",         "webApplicationFirewallConfiguration.ruleSetType"),
    ("WAFルールセットバージョン",     "webApplicationFirewallConfiguration.ruleSetVersion"),
    ("フロントエンドポート",          "frontendPorts[*].port"),
]

AGW_POOL_FIELDS = [
    ("プール名",                      "name"),
    ("IPアドレス",                    "backendAddresses[*].ipAddress"),
    ("FQDN",                          "backendAddresses[*].fqdn"),
]

AGW_HTTP_SETTINGS_FIELDS = [
    ("設定名",                        "name"),
    ("ポート",                        "port"),
    ("プロトコル",                    "protocol"),
    ("クッキーアフィニティ",          "cookieBasedAffinity"),
    ("タイムアウト(秒)",              "requestTimeout"),
    ("ホスト名(バックエンドから)",    "pickHostNameFromBackendAddress"),
    ("カスタムプローブ",              "probe.id"),
]

AGW_LISTENER_FIELDS = [
    ("リスナー名",                    "name"),
    ("プロトコル",                    "protocol"),
    ("フロントエンドポート",          "frontendPort.id"),
    ("ホスト名",                      "hostName"),
    ("SSL証明書",                     "sslCertificate.id"),
]

AGW_RULE_FIELDS = [
    ("ルール名",                      "name"),
    ("優先度",                        "priority"),
    ("ルールタイプ",                  "ruleType"),
    ("リスナー",                      "httpListener.id"),
    ("バックエンドプール",            "backendAddressPool.id"),
    ("バックエンドHTTP設定",          "backendHttpSettings.id"),
    ("URLパスマップ",                 "urlPathMap.id"),
    ("リダイレクト設定",              "redirectConfiguration.id"),
]

FIREWALL_FIELDS = [
    ("名前",                          "name"),
    ("リージョン",                    "location"),
    ("SKU名",                         "sku.name"),
    ("SKUティア",                     "sku.tier"),
    ("ファイアウォールポリシー",      "firewallPolicy.id"),
    ("脅威インテリジェンスモード",    "threatIntelMode"),
    ("IPアドレス設定名",              "ipConfigurations[*].name"),
    ("サブネット",                    "ipConfigurations[*].subnet.id"),
    ("パブリックIP",                  "ipConfigurations[*].publicIPAddress.id"),
    ("プライベートIPアドレス",        "ipConfigurations[*].privateIPAddress"),
    ("管理IP設定名",                  "managementIpConfiguration.name"),
    ("Availability Zones",            "zones"),
]

# ============================================================
# スタイル定義
# ============================================================

COLOR_HEADER_SECTION = "1F4E79"
COLOR_HEADER_COL     = "2E75B6"
COLOR_RESOURCE_ROW   = "E2EFDA"
COLOR_ODD_ROW        = "FFFFFF"
COLOR_EVEN_ROW       = "F2F7FB"
COLOR_NG_FILL        = "FFB3B3"   # 赤ハイライト（差分あり）
COLOR_OK_FILL        = "E2EFDA"   # 緑（一致）
COLOR_SUMMARY_NG     = "FF0000"
COLOR_SUMMARY_OK     = "70AD47"


def _side():
    return Side(style="thin", color="BFBFBF")

def thin_border():
    s = _side()
    return Border(left=s, right=s, top=s, bottom=s)

def cell_font(bold=False, color="000000", size=10):
    return Font(name="Arial", bold=bold, color=color, size=size)

def cell_fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def set_header(ws, row, col, value, bg=COLOR_HEADER_COL, fg="FFFFFF"):
    c = ws.cell(row=row, column=col, value=value)
    c.font = cell_font(bold=True, color=fg, size=10)
    c.fill = cell_fill(bg)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = thin_border()
    return c

def set_data(ws, row, col, value, bg=COLOR_ODD_ROW, bold=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = cell_font(bold=bold)
    c.fill = cell_fill(bg)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    c.border = thin_border()
    return c

def write_section_title(ws, row, title, ncols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=title)
    c.font = cell_font(bold=True, color="FFFFFF", size=12)
    c.fill = cell_fill(COLOR_HEADER_SECTION)
    c.alignment = Alignment(horizontal="left", vertical="center")
    c.border = thin_border()
    ws.row_dimensions[row].height = 22

def write_resource_header(ws, row, label, ncols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=label)
    c.font = cell_font(bold=True, size=10)
    c.fill = cell_fill(COLOR_RESOURCE_ROW)
    c.border = thin_border()
    ws.row_dimensions[row].height = 18

# ============================================================
# resolve（generate_param_sheet.py と同一）
# ============================================================

def resolve(obj, path):
    if obj is None:
        return ""
    if "[*]" in path:
        parts = path.split("[*].")
        base_path = parts[0]
        rest = parts[1] if len(parts) > 1 else None
        val = resolve(obj, base_path)
        if isinstance(val, list):
            if rest:
                items = [resolve(i, rest) for i in val if i]
                items = [str(i) for i in items if i != ""]
            else:
                items = [str(i) for i in val if i is not None]
            return ", ".join(items) if items else ""
        return ""
    keys = path.split(".")
    current = obj
    for k in keys:
        if current is None:
            return ""
        if isinstance(current, dict):
            current = current.get(k)
        elif isinstance(current, list):
            current = [item.get(k) for item in current if isinstance(item, dict)]
        else:
            return ""
    if isinstance(current, list):
        items = [str(i) for i in current if i is not None]
        return ", ".join(items) if items else ""
    if isinstance(current, bool):
        return "有効" if current else "無効"
    if current is None:
        return ""
    return str(current)

# ============================================================
# 設計書 Excel パーサー
# ============================================================

def _cell_val(row, col_idx):
    """0-indexed。値を文字列に正規化して返す。"""
    if col_idx >= len(row):
        return ""
    v = row[col_idx].value
    if v is None:
        return ""
    if isinstance(v, bool):
        return "有効" if v else "無効"
    return str(v).strip()

def _is_resource_header(text):
    """セルがリソースヘッダー行かどうか"""
    prefixes = (
        "VNet:", "Subnet:", "NSG:", "リソース:",
        "AGW:", "バックエンドプール:", "HTTP設定:", "リスナー:", "ルーティングルール:",
    )
    return any(text.startswith(p) for p in prefixes)

def parse_kvp_sheet(ws, value_col, apipath_col):
    """
    key-value形式シートをパース。
    value_col / apipath_col: 1-indexed 列番号
    戻り値: {resource_key: {api_path: design_value}}
    """
    result = {}
    current_resource = None

    for row in ws.iter_rows():
        first = _cell_val(row, 0)
        if not first:
            continue
        if _is_resource_header(first):
            current_resource = first
            result[current_resource] = {}
            continue
        # データ行判定（#列が数値）
        try:
            int(float(first))
        except (ValueError, TypeError):
            continue
        if current_resource is None:
            continue
        api_path = _cell_val(row, apipath_col - 1)
        design_val = _cell_val(row, value_col - 1)
        if api_path:
            result[current_resource][api_path] = design_val

    return result

def parse_nsg_sheet(ws):
    """
    NSGシートをパース（横並びルール形式）。
    戻り値: {nsg_name: [{col_name: value, ...}, ...]}
    """
    result = {}
    current_nsg = None
    header_row_idx = None
    col_map = {}  # col_index -> column_name

    for ri, row in enumerate(ws.iter_rows()):
        first = _cell_val(row, 0)

        # NSGヘッダー行
        if first.startswith("NSG:"):
            current_nsg = first[4:].strip()
            result[current_nsg] = []
            continue

        # ヘッダー行検出（"#" が先頭 — NSGリソース行より前に来るため current_nsg チェック不要）
        if first == "#":
            col_map = {}
            for ci, cell in enumerate(row):
                v = cell.value
                if v is not None:
                    col_map[ci] = str(v)
            header_row_idx = ri
            continue

        # データ行（#列が数値）
        if current_nsg is not None and col_map:
            try:
                int(float(first))
            except (ValueError, TypeError):
                continue
            rule = {}
            for ci, col_name in col_map.items():
                rule[col_name] = _cell_val(row, ci)
            result[current_nsg].append(rule)

    return result

# ============================================================
# 実環境データ → 突合用辞書生成
# ============================================================

def build_actual_kvp(resources, key_prefix, fields):
    """
    key-value形式リソースの突合辞書を構築。
    戻り値: {resource_key: {api_path: actual_value}}
    """
    result = {}
    for r in resources:
        name = r.get("name", r.get("id", "unknown"))
        key = f"{key_prefix}: {name}"
        result[key] = {}
        for _ui, api_path in fields:
            result[key][api_path] = resolve(r, api_path)
    return result

def build_actual_vnet(vnets):
    vnet_data = build_actual_kvp(vnets, "VNet", VNET_FIELDS)

    subnet_data = {}
    for vnet in vnets:
        vnet_name = vnet.get("name", "")
        for subnet in vnet.get("subnets", []):
            sname = subnet.get("name", "")
            key = f"Subnet: {vnet_name} / {sname}"
            subnet_data[key] = {}
            for _ui, api_path in SUBNET_FIELDS:
                subnet_data[key][api_path] = resolve(subnet, api_path)

    return vnet_data, subnet_data

def build_actual_nsg(nsgs):
    """
    戻り値: {nsg_name: [{col_name: value, ...}, ...]}
    """
    result = {}
    for nsg in nsgs:
        name = nsg.get("name", "")
        rules = []
        for idx, rule in enumerate(nsg.get("securityRules", [])):
            entry = {"#": str(idx + 1), "NSG名": name}
            entry["ルール名"]   = rule.get("name", "")
            entry["優先度"]     = str(rule.get("priority", ""))
            entry["方向"]       = rule.get("direction", "")
            entry["アクセス"]   = rule.get("access", "")
            entry["プロトコル"] = rule.get("protocol", "")
            entry["ソースIP"]   = rule.get("sourceAddressPrefix", "")
            entry["ソースPort"] = rule.get("sourcePortRange", "")
            entry["宛先IP"]     = rule.get("destinationAddressPrefix", "")
            entry["宛先Port"]   = rule.get("destinationPortRange", "")
            entry["説明"]       = rule.get("description", "") or ""
            rules.append(entry)
        result[name] = rules
    return result

def build_actual_agw(agws):
    """
    AGWの突合辞書を構築。
    戻り値: {resource_key: {api_path: actual_value}}
    resource_key 例:
      "AGW: agw-name"
      "バックエンドプール: agw-name / pool-name"
      "HTTP設定: agw-name / setting-name"
      "リスナー: agw-name / listener-name"
      "ルーティングルール: agw-name / rule-name"
    """
    result = {}
    for agw in agws:
        agw_name = agw.get("name", "")

        key = f"AGW: {agw_name}"
        result[key] = {p: resolve(agw, p) for _, p in AGW_TOP_FIELDS}

        for pool in agw.get("backendAddressPools", []):
            k = f"バックエンドプール: {agw_name} / {pool.get('name','')}"
            result[k] = {p: resolve(pool, p) for _, p in AGW_POOL_FIELDS}

        for setting in agw.get("backendHttpSettingsCollection", []):
            obj = setting.get("properties", setting)
            obj.setdefault("name", setting.get("name", ""))
            k = f"HTTP設定: {agw_name} / {setting.get('name','')}"
            result[k] = {p: resolve(obj, p) for _, p in AGW_HTTP_SETTINGS_FIELDS}

        for listener in agw.get("httpListeners", []):
            obj = listener.get("properties", listener)
            obj.setdefault("name", listener.get("name", ""))
            k = f"リスナー: {agw_name} / {listener.get('name','')}"
            result[k] = {p: resolve(obj, p) for _, p in AGW_LISTENER_FIELDS}

        for rule in agw.get("requestRoutingRules", []):
            obj = rule.get("properties", rule)
            obj.setdefault("name", rule.get("name", ""))
            k = f"ルーティングルール: {agw_name} / {rule.get('name','')}"
            result[k] = {p: resolve(obj, p) for _, p in AGW_RULE_FIELDS}

    return result

# ============================================================
# 差分判定
# ============================================================

def values_match(design_val, actual_val):
    """大文字小文字・前後空白を無視した比較。"""
    return design_val.strip().lower() == actual_val.strip().lower()

# ============================================================
# 差分シート生成
# ============================================================

NCOLS_KVP = 7   # # / UI項目名 / 設計値 / 実環境値 / 判定 / APIプロパティパス / 備考
NCOLS_SUB = 8   # # / VNet名 / UI項目名 / 設計値 / 実環境値 / 判定 / APIプロパティパス / 備考

KVP_HEADERS     = ["#", "UI項目名", "設計値", "実環境値", "判定", "APIプロパティパス", "備考"]
KVP_COL_WIDTHS  = [5, 35, 40, 40, 8, 50, 25]

SUB_HEADERS     = ["#", "VNet名", "UI項目名", "設計値", "実環境値", "判定", "APIプロパティパス", "備考"]
SUB_COL_WIDTHS  = [5, 25, 35, 38, 38, 8, 50, 25]

NSG_HEADERS     = ["#", "NSG名", "ルール名", "優先度", "方向", "アクセス",
                   "プロトコル", "ソースIP", "ソースPort", "宛先IP", "宛先Port", "説明", "判定", "備考"]
NSG_COL_WIDTHS  = [5, 25, 25, 10, 10, 10, 12, 25, 15, 25, 15, 30, 8, 25]


def _write_kvp_rows(ws, row, ui_name, api_path, design_val, actual_val, idx, vnet_name=None):
    """key-value形式の1データ行を書く。差分があれば赤ハイライト。"""
    is_ng = not values_match(design_val, actual_val)
    bg_data = COLOR_NG_FILL if is_ng else (COLOR_ODD_ROW if idx % 2 == 0 else COLOR_EVEN_ROW)
    judgment = "NG" if is_ng else "OK"
    j_color = COLOR_SUMMARY_NG if is_ng else COLOR_SUMMARY_OK

    if vnet_name is not None:
        set_data(ws, row, 1, idx + 1, bg_data)
        set_data(ws, row, 2, vnet_name, bg_data)
        set_data(ws, row, 3, ui_name, bg_data)
        set_data(ws, row, 4, design_val, bg_data)
        set_data(ws, row, 5, actual_val, bg_data)
        c = ws.cell(row=row, column=6, value=judgment)
        c.font = cell_font(bold=True, color=j_color)
        c.fill = cell_fill(bg_data)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border()
        p = set_data(ws, row, 7, api_path, bg_data)
        p.font = Font(name="Courier New", size=9, color="555555")
        set_data(ws, row, 8, "", bg_data)
    else:
        set_data(ws, row, 1, idx + 1, bg_data)
        set_data(ws, row, 2, ui_name, bg_data)
        set_data(ws, row, 3, design_val, bg_data)
        set_data(ws, row, 4, actual_val, bg_data)
        c = ws.cell(row=row, column=5, value=judgment)
        c.font = cell_font(bold=True, color=j_color)
        c.fill = cell_fill(bg_data)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border()
        p = set_data(ws, row, 6, api_path, bg_data)
        p.font = Font(name="Courier New", size=9, color="555555")
        set_data(ws, row, 7, "", bg_data)

    return is_ng


def create_vnet_diff_sheet(wb, design_vnet, design_subnet, actual_vnet, actual_subnet):
    ws = wb.create_sheet("VNet・Subnet")
    ws.freeze_panes = "A3"
    ng_count = 0

    # ---- VNet セクション ----
    write_section_title(ws, 1, "■ Virtual Network", NCOLS_KVP)
    for ci, (h, w) in enumerate(zip(KVP_HEADERS, KVP_COL_WIDTHS), 1):
        set_header(ws, 2, ci, h)
        ws.column_dimensions[get_column_letter(ci)].width = w
    row = 3

    all_vnet_keys = sorted(set(list(design_vnet.keys()) + list(actual_vnet.keys())))
    for vkey in all_vnet_keys:
        write_resource_header(ws, row, vkey, NCOLS_KVP)
        row += 1
        d_res = design_vnet.get(vkey, {})
        a_res = actual_vnet.get(vkey, {})
        all_paths = [p for _, p in VNET_FIELDS]
        ui_map = {p: u for u, p in VNET_FIELDS}
        for idx, api_path in enumerate(all_paths):
            d_val = d_res.get(api_path, "")
            a_val = a_res.get(api_path, "")
            is_ng = _write_kvp_rows(ws, row, ui_map[api_path], api_path, d_val, a_val, idx)
            if is_ng:
                ng_count += 1
            row += 1
        row += 1

    # ---- Subnet セクション ----
    write_section_title(ws, row, "■ Subnet", NCOLS_SUB)
    row += 1
    for ci, (h, w) in enumerate(zip(SUB_HEADERS, SUB_COL_WIDTHS), 1):
        set_header(ws, row, ci, h)
        ws.column_dimensions[get_column_letter(ci)].width = max(
            ws.column_dimensions[get_column_letter(ci)].width, w)
    row += 1

    all_sub_keys = sorted(set(list(design_subnet.keys()) + list(actual_subnet.keys())))
    for skey in all_sub_keys:
        # "Subnet: <vnet> / <subnet>" → vnet名を取り出す
        vnet_name = skey.replace("Subnet:", "").split("/")[0].strip() if "/" in skey else ""
        write_resource_header(ws, row, skey, NCOLS_SUB)
        row += 1
        d_res = design_subnet.get(skey, {})
        a_res = actual_subnet.get(skey, {})
        all_paths = [p for _, p in SUBNET_FIELDS]
        ui_map = {p: u for u, p in SUBNET_FIELDS}
        for idx, api_path in enumerate(all_paths):
            d_val = d_res.get(api_path, "")
            a_val = a_res.get(api_path, "")
            is_ng = _write_kvp_rows(ws, row, ui_map[api_path], api_path, d_val, a_val, idx,
                                    vnet_name=vnet_name)
            if is_ng:
                ng_count += 1
            row += 1
        row += 1

    return ws, ng_count


def create_nsg_diff_sheet(wb, design_nsg, actual_nsg):
    ws = wb.create_sheet("NSG")
    ws.freeze_panes = "A3"
    ng_count = 0
    ncols = len(NSG_HEADERS)

    write_section_title(ws, 1, "■ Network Security Group", ncols)
    for ci, (h, w) in enumerate(zip(NSG_HEADERS, NSG_COL_WIDTHS), 1):
        set_header(ws, 2, ci, h)
        ws.column_dimensions[get_column_letter(ci)].width = w
    row = 3

    all_nsg_names = sorted(set(list(design_nsg.keys()) + list(actual_nsg.keys())))
    for nsg_name in all_nsg_names:
        write_resource_header(ws, row, f"NSG: {nsg_name}", ncols)
        row += 1

        d_rules = design_nsg.get(nsg_name, [])
        a_rules = actual_nsg.get(nsg_name, [])

        # ルール名でマッチング
        a_map = {r.get("ルール名", ""): r for r in a_rules}

        for idx, d_rule in enumerate(d_rules):
            rule_name = d_rule.get("ルール名", "")
            a_rule = a_map.get(rule_name, {})

            compare_cols = ["ルール名", "優先度", "方向", "アクセス", "プロトコル",
                            "ソースIP", "ソースPort", "宛先IP", "宛先Port", "説明"]
            row_ng = any(
                not values_match(d_rule.get(c, ""), a_rule.get(c, ""))
                for c in compare_cols
            )
            if row_ng:
                ng_count += 1

            bg = COLOR_NG_FILL if row_ng else (COLOR_ODD_ROW if idx % 2 == 0 else COLOR_EVEN_ROW)
            judgment = "NG" if row_ng else "OK"
            j_color = COLOR_SUMMARY_NG if row_ng else COLOR_SUMMARY_OK

            vals = [
                idx + 1, nsg_name,
                d_rule.get("ルール名", ""), d_rule.get("優先度", ""),
                d_rule.get("方向", ""), d_rule.get("アクセス", ""),
                d_rule.get("プロトコル", ""), d_rule.get("ソースIP", ""),
                d_rule.get("ソースPort", ""), d_rule.get("宛先IP", ""),
                d_rule.get("宛先Port", ""), d_rule.get("説明", ""),
            ]
            for ci, v in enumerate(vals, 1):
                set_data(ws, row, ci, v, bg)
            c = ws.cell(row=row, column=13, value=judgment)
            c.font = cell_font(bold=True, color=j_color)
            c.fill = cell_fill(bg)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border()
            set_data(ws, row, 14, "", bg)
            row += 1

        # 実環境にしかないルールを追加表示
        d_rule_names = {r.get("ルール名", "") for r in d_rules}
        for idx2, a_rule in enumerate(a_rules):
            rname = a_rule.get("ルール名", "")
            if rname not in d_rule_names:
                bg = COLOR_NG_FILL
                ng_count += 1
                vals = [
                    len(d_rules) + idx2 + 1, nsg_name,
                    rname, a_rule.get("優先度", ""),
                    a_rule.get("方向", ""), a_rule.get("アクセス", ""),
                    a_rule.get("プロトコル", ""), a_rule.get("ソースIP", ""),
                    a_rule.get("ソースPort", ""), a_rule.get("宛先IP", ""),
                    a_rule.get("宛先Port", ""), a_rule.get("説明", ""),
                ]
                for ci, v in enumerate(vals, 1):
                    set_data(ws, row, ci, v, bg)
                c = ws.cell(row=row, column=13, value="NG(実環境のみ)")
                c.font = cell_font(bold=True, color=COLOR_SUMMARY_NG)
                c.fill = cell_fill(bg)
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.border = thin_border()
                set_data(ws, row, 14, "設計書にないルール", bg)
                row += 1

        row += 1

    return ws, ng_count


def create_generic_diff_sheet(wb, sheet_name, title, fields, design_data, actual_data):
    ws = wb.create_sheet(sheet_name)
    ws.freeze_panes = "A3"
    ng_count = 0

    write_section_title(ws, 1, title, NCOLS_KVP)
    for ci, (h, w) in enumerate(zip(KVP_HEADERS, KVP_COL_WIDTHS), 1):
        set_header(ws, 2, ci, h)
        ws.column_dimensions[get_column_letter(ci)].width = w
    row = 3

    all_keys = sorted(set(list(design_data.keys()) + list(actual_data.keys())))
    for rkey in all_keys:
        write_resource_header(ws, row, rkey, NCOLS_KVP)
        row += 1
        d_res = design_data.get(rkey, {})
        a_res = actual_data.get(rkey, {})
        all_paths = [p for _, p in fields]
        ui_map = {p: u for u, p in fields}
        for idx, api_path in enumerate(all_paths):
            d_val = d_res.get(api_path, "")
            a_val = a_res.get(api_path, "")
            is_ng = _write_kvp_rows(ws, row, ui_map[api_path], api_path, d_val, a_val, idx)
            if is_ng:
                ng_count += 1
            row += 1
        row += 1

    return ws, ng_count


def create_agw_diff_sheet(wb, design_agw, actual_agw):
    """
    AGW専用差分シート。サブセクション（プール/HTTP設定/リスナー/ルール）ごとに比較。
    design_agw / actual_agw: {resource_key: {api_path: value}}
    """
    ws = wb.create_sheet("AppGateway")
    ws.freeze_panes = "A3"
    ng_count = 0

    sections = [
        ("■ Application Gateway",   "AGW:",              AGW_TOP_FIELDS),
        ("■ バックエンドプール",     "バックエンドプール:", AGW_POOL_FIELDS),
        ("■ バックエンドHTTP設定",   "HTTP設定:",          AGW_HTTP_SETTINGS_FIELDS),
        ("■ HTTPリスナー",           "リスナー:",          AGW_LISTENER_FIELDS),
        ("■ ルーティングルール",     "ルーティングルール:", AGW_RULE_FIELDS),
    ]

    row = 1
    for section_title, key_prefix, fields in sections:
        write_section_title(ws, row, section_title, NCOLS_KVP)
        row += 1
        for ci, (h, w) in enumerate(zip(KVP_HEADERS, KVP_COL_WIDTHS), 1):
            set_header(ws, row, ci, h)
            ws.column_dimensions[get_column_letter(ci)].width = max(
                ws.column_dimensions[get_column_letter(ci)].width, w)
        row += 1

        all_keys = sorted(
            k for k in set(list(design_agw.keys()) + list(actual_agw.keys()))
            if k.startswith(key_prefix)
        )
        for rkey in all_keys:
            write_resource_header(ws, row, rkey, NCOLS_KVP)
            row += 1
            d_res = design_agw.get(rkey, {})
            a_res = actual_agw.get(rkey, {})
            all_paths = [p for _, p in fields]
            ui_map = {p: u for u, p in fields}
            for idx, api_path in enumerate(all_paths):
                d_val = d_res.get(api_path, "")
                a_val = a_res.get(api_path, "")
                is_ng = _write_kvp_rows(ws, row, ui_map[api_path], api_path, d_val, a_val, idx)
                if is_ng:
                    ng_count += 1
                row += 1
            row += 1

    return ws, ng_count


def create_summary_sheet(wb, rg_name, sheet_ng_counts):
    ws = wb.create_sheet("サマリー", 0)
    write_section_title(ws, 1, "Azure パラメータ突合チェック結果", 4)

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 30

    set_header(ws, 2, 1, "シート名")
    set_header(ws, 2, 2, "NG件数")
    set_header(ws, 2, 3, "判定")
    set_header(ws, 2, 4, "備考")

    total_ng = 0
    for idx, (sname, ng) in enumerate(sheet_ng_counts, 3):
        total_ng += ng
        bg = COLOR_NG_FILL if ng > 0 else COLOR_OK_FILL
        set_data(ws, idx, 1, sname, bg)
        set_data(ws, idx, 2, ng, bg)
        judgment = "NG" if ng > 0 else "OK"
        j_color = COLOR_SUMMARY_NG if ng > 0 else COLOR_SUMMARY_OK
        c = ws.cell(row=idx, column=3, value=judgment)
        c.font = cell_font(bold=True, color=j_color)
        c.fill = cell_fill(bg)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border()
        set_data(ws, idx, 4, "", bg)

    # 合計行
    total_row = len(sheet_ng_counts) + 3
    total_bg = COLOR_NG_FILL if total_ng > 0 else COLOR_OK_FILL
    set_data(ws, total_row, 1, "合計", total_bg, bold=True)
    set_data(ws, total_row, 2, total_ng, total_bg, bold=True)
    j = "NG" if total_ng > 0 else "OK"
    jc = COLOR_SUMMARY_NG if total_ng > 0 else COLOR_SUMMARY_OK
    c = ws.cell(row=total_row, column=3, value=j)
    c.font = cell_font(bold=True, color=jc)
    c.fill = cell_fill(total_bg)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = thin_border()
    set_data(ws, total_row, 4, "", total_bg)

    ws.cell(row=total_row + 2, column=1, value=f"対象RG: {rg_name}")
    ws.cell(row=total_row + 2, column=1).font = cell_font(bold=True)

    return ws

# ============================================================
# メイン
# ============================================================

def main():
    if len(sys.argv) < 3:
        print("Usage: python3 check_param_diff.py <design.xlsx> <azure_params.json> [output.xlsx]")
        sys.exit(1)

    design_file = sys.argv[1]
    json_file   = sys.argv[2]
    output_file = sys.argv[3] if len(sys.argv) > 3 else \
        json_file.replace(".json", "_diff.xlsx")

    # ---- 設計書ロード ----
    print(f"設計書読み込み中: {design_file}")
    design_wb = load_workbook(design_file, data_only=True)

    def get_sheet(name):
        return design_wb[name] if name in design_wb.sheetnames else None

    vnet_ws    = get_sheet("VNet・Subnet")
    nsg_ws     = get_sheet("NSG")
    sa_ws      = get_sheet("StorageAccount")
    kv_ws      = get_sheet("KeyVault")
    cae_ws     = get_sheet("ContainerAppsEnv")
    ca_ws      = get_sheet("ContainerApps")
    agw_ws     = get_sheet("AppGateway")
    fw_ws      = get_sheet("Firewall")

    # VNet(col3=設計値, col4=APIパス) / Subnet(col4=設計値, col5=APIパス)
    design_vnet   = parse_kvp_sheet(vnet_ws, value_col=3, apipath_col=4) if vnet_ws else {}
    design_subnet = {}
    if vnet_ws:
        # サブネットは6列構成: #/VNet名/UI項目名/設計値/APIパス/備考
        for k, v in parse_kvp_sheet(vnet_ws, value_col=4, apipath_col=5).items():
            if k.startswith("Subnet:"):
                design_subnet[k] = v
            elif k.startswith("VNet:") and k not in design_vnet:
                design_vnet[k] = v
        # VNetキーのみ再抽出（value_col=3 で取った結果）
        design_vnet = {k: v for k, v in design_vnet.items() if k.startswith("VNet:")}

    design_nsg = parse_nsg_sheet(nsg_ws) if nsg_ws else {}
    design_sa  = parse_kvp_sheet(sa_ws,  value_col=3, apipath_col=4) if sa_ws  else {}
    design_kv  = parse_kvp_sheet(kv_ws,  value_col=3, apipath_col=4) if kv_ws  else {}
    design_cae = parse_kvp_sheet(cae_ws, value_col=3, apipath_col=4) if cae_ws else {}
    design_ca  = parse_kvp_sheet(ca_ws,  value_col=3, apipath_col=4) if ca_ws  else {}
    design_agw = parse_kvp_sheet(agw_ws, value_col=3, apipath_col=4) if agw_ws else {}
    design_fw  = parse_kvp_sheet(fw_ws,  value_col=3, apipath_col=4) if fw_ws  else {}

    # ---- 実環境データロード ----
    print(f"実環境データ読み込み中: {json_file}")
    with open(json_file, encoding="utf-8") as f:
        data = json.load(f)

    actual_vnet, actual_subnet = build_actual_vnet(data.get("vnets", []))
    actual_nsg = build_actual_nsg(data.get("nsgs", []))
    actual_sa  = build_actual_kvp(data.get("storageAccounts", []),     "リソース", STORAGE_FIELDS)
    actual_kv  = build_actual_kvp(data.get("keyVaults", []),           "リソース", KEYVAULT_FIELDS)
    actual_cae = build_actual_kvp(data.get("containerAppEnvs", []),    "リソース", CAE_FIELDS)
    actual_ca  = build_actual_kvp(data.get("containerApps", []),       "リソース", CA_FIELDS)
    actual_agw = build_actual_agw(data.get("applicationGateways", []))
    actual_fw  = build_actual_kvp(data.get("firewalls", []),           "リソース", FIREWALL_FIELDS)

    # ---- 出力 Excel 生成 ----
    print("差分チェック中...")
    wb = Workbook()
    wb.remove(wb.active)

    _, ng_vnet = create_vnet_diff_sheet(wb, design_vnet, design_subnet, actual_vnet, actual_subnet)
    _, ng_nsg  = create_nsg_diff_sheet(wb, design_nsg, actual_nsg)
    _, ng_sa   = create_generic_diff_sheet(wb, "StorageAccount",   "■ Storage Account",            STORAGE_FIELDS,  design_sa,  actual_sa)
    _, ng_kv   = create_generic_diff_sheet(wb, "KeyVault",         "■ Key Vault",                  KEYVAULT_FIELDS, design_kv,  actual_kv)
    _, ng_cae  = create_generic_diff_sheet(wb, "ContainerAppsEnv", "■ Container Apps Environment", CAE_FIELDS,      design_cae, actual_cae)
    _, ng_ca   = create_generic_diff_sheet(wb, "ContainerApps",    "■ Container Apps",             CA_FIELDS,       design_ca,  actual_ca)
    _, ng_agw  = create_agw_diff_sheet(wb, design_agw, actual_agw)
    _, ng_fw   = create_generic_diff_sheet(wb, "Firewall",         "■ Azure Firewall",             FIREWALL_FIELDS, design_fw,  actual_fw)

    sheet_ng_counts = [
        ("VNet・Subnet",     ng_vnet),
        ("NSG",              ng_nsg),
        ("StorageAccount",   ng_sa),
        ("KeyVault",         ng_kv),
        ("ContainerAppsEnv", ng_cae),
        ("ContainerApps",    ng_ca),
        ("AppGateway",       ng_agw),
        ("Firewall",         ng_fw),
    ]
    create_summary_sheet(wb, data.get("resourceGroup", ""), sheet_ng_counts)

    wb.save(output_file)
    total_ng = sum(n for _, n in sheet_ng_counts)
    print(f"\n突合チェック完了: {output_file}")
    print(f"{'シート名':<22} {'NG件数':>6}")
    print("-" * 30)
    for sname, ng in sheet_ng_counts:
        print(f"{sname:<22} {ng:>6}")
    print("-" * 30)
    print(f"{'合計':<22} {total_ng:>6}")
    if total_ng > 0:
        print("\n差分あり: 赤ハイライト行を確認してください。")
    else:
        print("\n差分なし: 全項目一致。")


if __name__ == "__main__":
    main()
